"""SQLite training registry, Excel review, LASSO and skeleton LSTM.

Run ``python -m ddc.training --help``. Pose detection is never repeated here.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import tempfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

from .dancer_features import load_labels
from .feature_inventory import collect_feature_inventory
from .feature_registry import load_merges, review_is_current
from .lasso_training import fit_lasso, score
from .skeleton_lstm import middle_sequences, train_lstm, predict_lstm, select_sequences, INPUT_VERSION, STEPS
from .training_features import FEATURES, FEATURE_IDS, PREDICTOR_IDS
from .training_store import DATABASE, LEGACY_DATABASE, TrainingStore, is_armed

WORKBOOK = "lasso_lstm_registry.xlsx"
COHORT_TARGET = 15


def _write(path, data):
    # Atomic replacement: readers see either complete version, never partial JSON.
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, allow_nan=False)
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def _write_labels(path, labels):
    fd, temp = tempfile.mkstemp(dir=path.parent, suffix=".csv.tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["track_id", "label"])
            writer.writerows((tid, "dancer" if label else "non_dancer")
                             for tid, label in sorted(labels.items()))
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def _file_digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


@contextmanager
def _lock(root):
    from filelock import FileLock
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    with FileLock(str(root / ".training.lock"), timeout=30):
        yield


def _validated_run(run, *, existing=False):
    """Only a complete, current human review can claim a cohort slot."""
    run = Path(run)
    has_review = (run / "training_review.json").exists() or (run / "prediction_validation.json").exists()
    if not has_review and not existing:
        return False
    try:
        meta = json.loads((run / "meta.json").read_text(encoding="utf-8"))
        if meta.get("frames", 0) <= 0 or not (run / "tracking.jsonl").stat().st_size:
            return False
        current = review_is_current(run)
        sig = signature(run)
        stored = TrainingStore(run.parent).reviews(run.name, sig) if existing else {}
        if has_review and not current and not stored:
            return False
        labels = (load_labels(run / "labels.csv")
                  if (run / "labels.csv").exists() and current else stored)
        if not labels:
            return False
        rows = (_feature_rows(run, load_merges(run), sig) if existing
                else collect_feature_inventory(run / "tracking.jsonl", load_merges(run)))
        return bool(rows) and {row["track_id"] for row in rows} == set(labels)
    except (OSError, ValueError, KeyError, TypeError):
        return False


def _normalize_state(root, state):
    """Upgrade older cohorts to fifteen validated videos."""
    if state.get("cohort_version", 1) < 2:
        state["target"] = COHORT_TARGET
        state["runs"] = [name for name in state["runs"] if _validated_run(root / name, existing=True)]
        state["membership_closed"] = len(state["runs"]) >= COHORT_TARGET
        state["frozen"] = False
        state["cohort_version"] = 2
    state.setdefault("membership_closed", len(state["runs"]) >= state["target"])
    state.setdefault("frozen", False)
    return state


def arm(runs_dir, target=COHORT_TARGET):
    if target != COHORT_TARGET:
        raise ValueError("The training cohort is fixed at fifteen validated videos")
    root = Path(runs_dir)
    with _lock(root):
        store = TrainingStore(root)
        state = store.state()
        if state is not None:
            state = _normalize_state(root, state)
            store.save_state(state)
            return state
        state = {"created_at": datetime.now(timezone.utc).isoformat(), "target": target,
                 "existing_runs": sorted(p.name for p in root.iterdir() if p.is_dir()),
                 "runs": [], "status": "collecting", "target_label": "dancer vs non_dancer",
                 "cohort_version": 2, "membership_closed": False, "frozen": False}
        store.save_state(state)
        return state


def register_run(runs_dir, run_name):
    """Enroll a new run only after every merged person has a current review."""
    root = Path(runs_dir)
    with _lock(root):
        store = TrainingStore(root)
        state = store.state()
        if state is None:
            return None
        state = _normalize_state(root, state)
        if run_name in state["existing_runs"] or run_name in state["runs"]:
            return state
        if state["membership_closed"] or state["frozen"] or len(state["runs"]) >= state["target"]:
            return state
        run = root / run_name
        if not _validated_run(run):
            return state
        meta = json.loads((run / "meta.json").read_text(encoding="utf-8"))
        state["runs"].append(run_name)
        meta.update(training_run=True, training_index=len(state["runs"]))
        _write(run / "meta.json", meta)
        state["membership_closed"] = len(state["runs"]) >= state["target"]
        store.save_state(state)
        return state


def signature(run_dir, merges=None):
    from .tracking import TRACKING_VERSION
    from .appearance import source_stamp
    from .identity_learning import prediction_stamp
    digest = hashlib.sha256()
    digest.update(f"tracking-v{TRACKING_VERSION}".encode())
    digest.update(json.dumps(source_stamp(Path(run_dir) / "tracking.jsonl")).encode())
    identity = prediction_stamp(Path(run_dir) / "tracking.jsonl")
    if identity is not None:
        digest.update(identity.encode())
    with (Path(run_dir) / "tracking.jsonl").open("rb") as f:
        for part in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(part)
    digest.update(json.dumps(load_merges(Path(run_dir)) if merges is None else merges).encode())
    return digest.hexdigest()


def _discover_runs(root, state):
    """Claim only reviewed runs; never extend a closed fifteen-run cohort."""
    if state["membership_closed"] or state["frozen"]:
        return
    missing = [name for name in state["runs"] if not (root / name / "tracking.jsonl").exists()]
    if missing:
        state["missing_runs"] = sorted(set(state.get("missing_runs", [])) | set(missing))
        state["runs"] = [name for name in state["runs"] if name not in missing]
    excluded = set(state["existing_runs"]) | set(state["runs"])
    for run in sorted(root.iterdir(), key=lambda p: p.name):
        if len(state["runs"]) >= state["target"]:
            break
        if not run.is_dir() or run.name in excluded:
            continue
        if not _validated_run(run):
            continue
        meta = json.loads((run / "meta.json").read_text(encoding="utf-8"))
        state["runs"].append(run.name)
        excluded.add(run.name)
        meta.update(training_run=True, training_index=len(state["runs"]))
        _write(run / "meta.json", meta)
    for index, name in enumerate(state["runs"], 1):
        path = root / name / "meta.json"
        meta = json.loads(path.read_text(encoding="utf-8"))
        if meta.get("training_index") != index:
            meta["training_index"] = index
            _write(path, meta)
    state["membership_closed"] = len(state["runs"]) >= state["target"]


def _feature_rows(run, merges, sig):
    store = TrainingStore(run.parent)
    rows = store.features(run.name, sig)
    if rows is not None:
        return rows
    cache = run / "training/features.json"
    if cache.exists():
        saved = json.loads(cache.read_text(encoding="utf-8"))
        if saved.get("signature") == sig:
            return store.features(run.name, sig, saved["tracks"])
    rows = collect_feature_inventory(run / "tracking.jsonl", merges)
    for row in rows:
        row["values"] = {fid: (float(row["values"][fid]) if isinstance(row["values"].get(fid),
                          (int, float)) and np.isfinite(row["values"][fid]) else None)
                          for fid in FEATURE_IDS}
    return store.features(run.name, sig, rows)


def _sequences(run, merges, sig):
    folder = run / "training"
    metadata, data = folder / "feature_preprocessing.json", folder / "middle_5s_features.npz"
    if metadata.exists() and data.exists():
        crop = json.loads(metadata.read_text(encoding="utf-8"))
        if (crop.get("signature") == sig and crop.get("input_version") == INPUT_VERSION
                and crop.get("feature_ids") == list(PREDICTOR_IDS)):
            with np.load(data, allow_pickle=False) as saved:
                tracks = {int(tid): {"sequence": seq, "coverage": float(coverage),
                                    "feature_ids": list(PREDICTOR_IDS)}
                          for tid, seq, coverage in zip(saved["track_ids"], saved["x"], saved["coverage"])}
            return tracks, crop
    tracks, crop = middle_sequences(run, merges)
    folder.mkdir(exist_ok=True)
    np.savez_compressed(data, x=np.stack([s["sequence"] for s in tracks.values()]) if tracks
                       else np.empty((0, STEPS, len(PREDICTOR_IDS)), np.float32),
                       track_ids=np.asarray(list(tracks)), coverage=np.asarray([s["coverage"] for s in tracks.values()]))
    _write(metadata, {**crop, "signature": sig})
    return tracks, crop


def import_reviews(runs_dir, workbook=None):
    """Validate the entire import before writing labels; blank labels stay unreviewed."""
    from openpyxl import load_workbook
    root = Path(runs_dir)
    path = Path(workbook) if workbook else root / WORKBOOK
    if not path.exists():
        return 0
    store = TrainingStore(root)
    state = store.state()
    if state is None:
        raise ValueError("Arm the fifteen-run cohort first")
    if state.get("frozen"):
        raise ValueError("The first fifteen-run model is frozen; workbook reviews cannot change training")
    pending = []
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for index, run_name in enumerate(state["runs"], 1):
            sheet_name = f"Run {index:02d}"
            if sheet_name not in wb.sheetnames:
                continue  # A just-enrolled run is not in the preceding export yet.
            ws = wb[sheet_name]
            sig = signature(root / run_name)
            if ws["B2"].value != run_name or ws["B3"].value != sig:
                # Merging tracks invalidates worksheet ids. Never attach old reviews to new people.
                if workbook and ws["B2"].value is not None:
                    raise ValueError(f"{sheet_name}: stale run or pose signature; download the refreshed workbook")
                continue
            label_path = root / run_name / "labels.csv"
            if label_path.exists() and label_path.stat().st_mtime_ns > path.stat().st_mtime_ns:
                continue  # A newer UI review takes precedence over an older Excel export.
            headers = [ws.cell(6, c).value for c in range(1, 7)]
            if (headers[:3] != ["Track id", "Frames present", "Predicted label"]
                    or headers[3] not in ("LASSO score", "Predictor score")
                    or headers[4] not in ("Manual label", "Validation label")
                    or headers[5] != "Review status"):
                raise ValueError(f"{sheet_name}: review column headings have changed")
            valid_ids = {r["track_id"] for r in _feature_rows(root / run_name,
                         load_merges(root / run_name), sig)}
            labels, seen = {}, set()
            for row in ws.iter_rows(min_row=7, values_only=True):
                if row[0] is None:
                    break
                tid, label = row[0], row[4]
                if not isinstance(tid, int) or tid not in valid_ids or tid in seen:
                    raise ValueError(f"{sheet_name}: unknown or duplicate track id {tid}")
                seen.add(tid)
                if label in (None, ""):
                    continue
                if label not in ("dancer", "non_dancer"):
                    raise ValueError(f"{sheet_name}, track {tid}: label must be dancer or non_dancer")
                labels[tid] = int(label == "dancer")
            if seen != valid_ids:
                raise ValueError(f"{sheet_name}: review table is missing track rows")
            if seen:
                pending.append((run_name, labels))
    finally:
        wb.close()
    store.import_reviews([(name, signature(root / name), labels,
                           _file_digest(root / name / "labels.csv")) for name, labels in pending])
    for run_name, labels in pending:
        run = root / run_name
        _write_labels(run / "labels.csv", labels)
        _write(run / "training_review.json", {"signature": signature(run),
               "reviewed_at": datetime.now(timezone.utc).isoformat(), "source": "Excel"})
        store.ingest_reviews(run_name, signature(run), labels, _file_digest(run / "labels.csv"))
    return sum(len(labels) for _, labels in pending)


def export_workbook(root, payload):
    """Use the bundled artifact tool when present, or configured Node/package runtime."""
    root = Path(root)
    transfer = root / "training_workbook_payload.json"
    _write(transfer, payload)
    try:
        bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node"
        node = os.environ.get("DDC_NODE") or (str(bundled / "bin/node.exe")
                 if (bundled / "bin/node.exe").exists() else shutil.which("node"))
        modules = os.environ.get("DDC_ARTIFACT_MODULE_ROOT") or str(bundled / "node_modules")
        if not node:
            raise RuntimeError("Excel export requires Node.js and @oai/artifact-tool")
        builder = Path(__file__).resolve().parent.parent / "scripts/build_training_workbook.mjs"
        done = subprocess.run([node, str(builder), str(root.resolve()), modules],
                              capture_output=True, text=True, encoding="utf-8", timeout=180)
        if done.returncode:
            raise RuntimeError(f"Excel export failed: {done.stderr[-1500:]}")
        return root / WORKBOOK
    finally:
        transfer.unlink(missing_ok=True)


def _export_database(root, store):
    """Excel failure leaves the committed snapshot pending for the next attempt."""
    try:
        path = export_workbook(root, store.payload())
    except (OSError, ValueError, RuntimeError, ImportError, subprocess.TimeoutExpired) as exc:
        store.sync_status(str(exc))
        return f"Excel update pending: {exc}. Close the workbook and run python -m ddc.training export."
    store.sync_status("")
    return ""


def export_database(runs_dir):
    """Regenerate Excel from the saved DB snapshot without fitting either model."""
    root = Path(runs_dir)
    with _lock(root):
        store = TrainingStore(root)
        warning = _export_database(root, store)
        if warning:
            raise RuntimeError(warning)
        return root / WORKBOOK


def refresh(runs_dir, train=True, epochs=60, import_excel=True, export=True, review_workbook=None):
    root = Path(runs_dir)
    with _lock(root):
        store = TrainingStore(root)
        state = store.state()
        if state is None:
            raise ValueError("Arm the fifteen-run cohort first")
        state = _normalize_state(root, state)
        if state["frozen"]:
            state["excel_sync"] = store.sync_status()
            return state
        _discover_runs(root, state)
        store.save_state(state)
        if import_excel:
            import_reviews(root, review_workbook)
        runs, records, problems, samples, inference = [], [], [], [], []
        for index, name in enumerate(state["runs"], 1):
            run = root / name
            if not _validated_run(run, existing=True):
                problems.append(f"{name}: tracking identities or reviewed labels changed; validate every track again")
                runs.append({"index": index, "run": name, "signature": "", "tracks": [],
                             "crop": {"error": "Awaiting a current complete review"}})
                continue
            merges = load_merges(run)
            sig = signature(run)
            try:
                rows = _feature_rows(run, merges, sig)
            except ValueError as exc:
                rows = []
                problems.append(f"{name}: {exc}")
            # Explicit UI/CSV inputs are ingested before fitting. Database reviews
            # survive a missing compatibility CSV, and are keyed by pose identity.
            labels = (store.ingest_reviews(name, sig, load_labels(run / "labels.csv"),
                                          _file_digest(run / "labels.csv"))
                      if (run / "labels.csv").exists() and review_is_current(run)
                      else store.reviews(name, sig))
            if not review_is_current(run) and not labels:
                problems.append(f"{name}: tracking identities changed; labels must be reviewed again")
            # Store only finite values in the JSON/Excel boundary.
            for row in rows:
                row["manual_label"] = ("dancer" if labels[row["track_id"]] else "non_dancer") \
                    if row["track_id"] in labels else ""
                if row["track_id"] in labels:
                    records.append({**row, "run": name, "label": labels[row["track_id"]]})
            if not rows or any(r["track_id"] not in labels for r in rows):
                problems.append(f"{name}: review every track's dancer/non_dancer label")
            try:
                sequences, crop = _sequences(run, merges, sig)
                for row in rows:
                    tid = row["track_id"]
                    if tid in sequences:
                        inference.append({**sequences[tid], "run": name, "track_id": tid})
                    if tid not in sequences:
                        problems.append(f"{name}, track {tid}: insufficient middle-window coverage or supported feature windows")
                    elif tid in labels:
                        samples.append({**sequences[tid], "run": name, "track_id": tid,
                                        "label": labels[tid], "crop": crop})
            except ValueError as exc:
                crop = {"error": str(exc)}
                problems.append(f"{name}: {exc}")
            runs.append({"index": index, "run": name, "signature": sig, "tracks": rows, "crop": crop})
        model = None
        try:
            model = fit_lasso(records)
        except ValueError as exc:
            problems.append(str(exc))
        if model:
            candidate_keys = {(s["run"], s["track_id"]) for s in inference}
            samples = select_sequences(samples, model["feature_ids"])
            inference = select_sequences(inference, model["feature_ids"])
            usable_keys = {(s["run"], s["track_id"]) for s in inference}
            problems.extend(f"{name}, track {tid}: insufficient measured values for the five selected features"
                            for name, tid in sorted(candidate_keys - usable_keys))
        else:
            samples, inference = [], []
        # Preserve the prediction actually reviewed in Dancers. A fresh model
        # score could change after this run joins the rolling training set.
        for run in runs:
            for row in run["tracks"]:
                row["prediction"] = None
                row["predictor_method"] = "Unavailable"
                row["review_action"] = ""
            folder = root / run["run"]
            saved_path = folder / "prediction_validation.json"
            saved = json.loads(saved_path.read_text()) if saved_path.exists() else {}
            review_path = folder / "training_review.json"
            review = json.loads(review_path.read_text()) if review_path.exists() else {}
            for row in run["tracks"]:
                key = str(row["track_id"])
                # A later model refit must not rewrite the prediction that was reviewed.
                if (row["manual_label"] and review.get("signature") == run["signature"]
                        and saved.get("predictor_labels", {}).get(key)):
                    row["prediction"] = {"predicted_label": saved["predictor_labels"][key],
                        "score": saved.get("predictor_probabilities", {}).get(key)}
                    row["predictor_method"] = "Saved Dancers prediction"
                row["review_action"] = saved.get("review_action", "Saved review") if row["manual_label"] else ""
                predicted = (row["prediction"] or {}).get("predicted_label")
                row["validation_status"] = ("Unreviewed" if not row["manual_label"] else
                    "Reviewed without prediction" if not predicted else
                    "Validated correct" if predicted == row["manual_label"] else "Corrected")
        run_signatures = {r["run"]: r["signature"] for r in runs}
        fingerprint = hashlib.sha256(json.dumps({
                   "samples": [(s["run"], s["track_id"], s["label"], run_signatures[s["run"]]) for s in samples],
                   "model": model, "epochs": epochs, "input_version": INPUT_VERSION}, sort_keys=True).encode()).hexdigest()
        report, checkpoint = store.model(fingerprint)
        report_path = root / "training/feature_lstm_report.json"
        if report is None and report_path.exists():
            legacy = json.loads(report_path.read_text())
            checkpoint_path = root / "training/feature_lstm.pt"
            if legacy.get("fingerprint") == fingerprint and checkpoint_path.exists():
                report, checkpoint = legacy, checkpoint_path.read_bytes()
        # Unreviewed/new clips do not stop updates from the reviewed usable clips.
        ready = model is not None and {s["label"] for s in samples} == {0, 1}
        contributing = {s["run"] for s in samples}
        complete = (state["membership_closed"] and len(runs) == state["target"]
                    and contributing == {r["run"] for r in runs})
        if state["membership_closed"] and not complete:
            absent = sorted({r["run"] for r in runs} - contributing)
            problems.append("The fixed fifteen-run cohort needs usable reviewed LSTM sequences from: "
                            + ", ".join(absent))
        training_failed = False
        if ready and train and (not checkpoint or not report or report.get("fingerprint") != fingerprint):
            state.update(status="training", lstm_current=False)
            store.save_state(state)
            try:
                report = train_lstm(samples, root / "training", model["feature_ids"], fingerprint, epochs)
                checkpoint = (root / "training/feature_lstm.pt").read_bytes()
            except (ImportError, OSError, RuntimeError, ValueError) as exc:
                problems.append(f"LSTM training failed: {exc}")
                training_failed = True
                ready = False
        current = bool(ready and checkpoint and report and report.get("fingerprint") == fingerprint)
        predictions = {}
        prediction_failed = False
        if current:
            try:
                predictions = predict_lstm(inference, checkpoint, fingerprint)
            except (ImportError, OSError, RuntimeError, ValueError) as exc:
                problems.append(f"LSTM prediction failed: {exc}")
                prediction_failed = True
                current = False
        trained_keys = {(s["run"], s["track_id"]) for s in samples}
        usable_keys = {(s["run"], s["track_id"]) for s in inference}
        for run in runs:
            for row in run["tracks"]:
                key = (run["run"], row["track_id"])
                result = predictions.get(key)
                row["lstm_prediction"] = ({**result, "fingerprint": fingerprint,
                    "in_training": key in trained_keys} if result and current else None)
                if result and current:
                    row["lstm_status"] = "Training fit" if key in trained_keys else "Unreviewed prediction"
                else:
                    row["lstm_status"] = ("No usable selected-feature sequence" if key not in usable_keys
                                          else "Awaiting current LSTM model")
        if current:
            status = "trained" if complete else "rolling_trained"
        elif prediction_failed:
            status = "prediction_failed"
        elif training_failed:
            status = "training_failed"
        elif ready:
            status = "ready_to_train"
        else:
            status = "collecting" if len(runs) < state["target"] else "awaiting_review_or_data"
        state.update(status=status, problems=problems, reviewed_tracks=len(records),
                     selected_features=model["feature_ids"] if model else [],
                     lstm_current=current, updated_at=datetime.now(timezone.utc).isoformat())
        state["frozen"] = bool(current and complete)
        state["lstm_runs"] = report["runs"] if current else []
        state["lstm_samples"] = len(samples)
        payload = {"state": state, "features": [s.__dict__ for s in FEATURES],
                   "model": model, "runs": runs, "lstm": report if current else None}
        store.save_snapshot(payload, fingerprint, checkpoint if current else None)
        if export:
            _export_database(root, store)
        state["excel_sync"] = store.sync_status()
        return state


def predict_run(runs_dir, run_name, merge_groups=None):
    """Score saved merged identities with the latest fitted LASSO and LSTM.

    This path only reads the SQLite model. A run outside the fixed training
    cohort cannot add examples, review rows, or model versions by being opened.
    """
    root = Path(runs_dir)
    if not (root / DATABASE).exists() and not (root / LEGACY_DATABASE).exists():
        raise ValueError("No saved LASSO/LSTM training model is available yet")
    store = TrainingStore(root)
    # A newly reviewed cohort run may change state before its next refresh has
    # committed a snapshot. The last saved fit remains usable for inference.
    with store.connect() as db:
        row = db.execute("SELECT payload_json FROM snapshots WHERE id=(SELECT CAST(value AS INTEGER) "
                         "FROM settings WHERE key='latest_snapshot')").fetchone()
    if row is None:
        raise ValueError("No saved LASSO/LSTM training model is available yet")
    payload = json.loads(row[0])
    model = payload.get("model")
    if not model:
        raise ValueError("Five LASSO features need validated dancer and non-dancer runs first")
    run = root / run_name
    merges = load_merges(run) if merge_groups is None else merge_groups
    rows = collect_feature_inventory(run / "tracking.jsonl", merges)
    if not rows:
        raise ValueError("No merged person tracks have measurable pose features")

    report = payload.get("lstm") or {}
    fingerprint = report.get("fingerprint")
    lstm_predictions = {}
    if fingerprint:
        _, checkpoint = store.model(fingerprint)
        if checkpoint:
            try:
                sequences, _ = middle_sequences(run, merges)
                candidates = [{**sequence, "run": run_name, "track_id": tid}
                              for tid, sequence in sequences.items()]
                selected = select_sequences(candidates, model["feature_ids"])
                if selected:
                    lstm_predictions = predict_lstm(selected, checkpoint, fingerprint)
            except (ImportError, OSError, RuntimeError, ValueError):
                pass  # Short or sparse clips can still use the saved LASSO fit.

    specs = {feature.id: feature for feature in FEATURES}
    selected_features = [{"rank": index + 1, "feature_id": fid,
                          "feature_name": specs[fid].name, "category": specs[fid].category,
                          "standardized_coefficient": model["coefficient"][index],
                          "ranking_weight": model["ranking_coefficients"][index],
                          "absolute_weight": abs(model["coefficient"][index]),
                          "direction": ("more dancer-like" if model["coefficient"][index] >= 0
                                        else "more non-dancer-like")}
                         for index, fid in enumerate(model["feature_ids"])]
    predictions = []
    for row in rows:
        linear = score(row, model)
        temporal = lstm_predictions.get((run_name, row["track_id"]))
        predictors = []
        for index, item in enumerate(selected_features):
            raw = row["values"].get(item["feature_id"])
            observed = (float(raw) if isinstance(raw, (int, float)) and np.isfinite(raw)
                        else None)
            effective = observed if observed is not None else model["medians"][index]
            standardized = (effective - model["mean"][index]) / model["std"][index]
            predictors.append({**item, "raw_value": observed,
                               "standardized_value": float(standardized),
                               "score_contribution": linear["contributions"][index]})
        predictions.append({
            "track_id": int(row["track_id"]), "frames_present": int(row["frames_present"]),
            "predicted_label": temporal["predicted_label"] if temporal else linear["predicted_label"],
            "dancer_probability": temporal["probability"] if temporal else None,
            "lstm_probability": temporal["probability"] if temporal else None,
            "lstm_label": temporal["predicted_label"] if temporal else None,
            "lasso_score": linear["score"], "lasso_label": linear["predicted_label"],
            "method": "LSTM probability" if temporal else "LASSO linear score",
            "top_predictors": predictors,
        })
    summary = {"feature_ids": list(model["feature_ids"]),
               "selected_features": selected_features,
               "reviewed_tracks": model["reviewed_tracks"],
               "labelled_tracks": model["reviewed_tracks"],
               "runs": list(model["runs"]),
               "fingerprint": fingerprint,
               "prediction_method": "LSTM probability with LASSO linear-score fallback"}
    return summary, predictions


def refresh_after_review(runs_dir):
    """Application hook: a workbook/model error must not discard a saved review."""
    try:
        if not is_armed(runs_dir):
            return ""
        state = refresh(runs_dir, import_excel=False)
        note = f" Training cohort: {len(state['runs'])}/{state['target']} videos, {state['status']}."
        sync = state.get("excel_sync") or {}
        if sync.get("status") == "pending":
            note += f" Excel update pending: {sync.get('error', '')}. Close the workbook and retry export."
        return note
    except (OSError, ValueError, RuntimeError, ImportError, sqlite3.Error, subprocess.TimeoutExpired) as exc:
        return f" Training refresh needs attention: {exc}"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["arm", "refresh", "status", "import-review", "export", "backup"])
    parser.add_argument("--runs-dir", type=Path, default=Path(os.environ.get("DDC_RUNS_DIR", "runs")))
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--no-train", action="store_true")
    parser.add_argument("--epochs", type=int, default=60)
    parser.add_argument("--output", type=Path, help="New destination file for a database backup")
    args = parser.parse_args()
    if args.command == "arm":
        arm(args.runs_dir)
        result = refresh(args.runs_dir, train=False)
    elif args.command == "status":
        store = TrainingStore(args.runs_dir)
        result = {"state": store.state(), "excel_sync": store.sync_status(), "database": str(store.path)}
    elif args.command == "export":
        result = {"workbook": str(export_database(args.runs_dir))}
    elif args.command == "backup":
        if not args.output:
            parser.error("backup requires --output")
        result = {"backup": str(TrainingStore(args.runs_dir).backup(args.output))}
    else:
        if args.command == "import-review":
            with _lock(args.runs_dir):
                import_reviews(args.runs_dir, args.workbook)
        result = refresh(args.runs_dir, train=not args.no_train, epochs=args.epochs,
                         import_excel=args.command != "import-review")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
