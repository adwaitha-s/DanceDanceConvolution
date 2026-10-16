import json

import numpy as np
import pytest

from ddc.dancer_features import write_synthetic_demo
from ddc.lasso_training import fit_lasso, score
from ddc.skeleton_lstm import make_model, middle_sequences, train_lstm
from ddc.training import arm, import_reviews, refresh, register_run, signature, replace_run
from ddc.training_features import FEATURES, PREDICTOR_IDS


def run(root, name, frames=180):
    folder = root / name
    folder.mkdir()
    write_synthetic_demo(folder / "tracking.jsonl", folder / "labels.csv", frames=frames)
    (folder / "meta.json").write_text(json.dumps({"frames": frames, "fps": 30, "duration_s": frames/30}))
    return folder


def test_catalog_has_exactly_40_across_requested_categories():
    assert len(FEATURES) == len({s.id for s in FEATURES}) == 40
    assert len({s.category for s in FEATURES}) == 5
    assert len(PREDICTOR_IDS) == 35
    assert not any(s.predictor_eligible for s in FEATURES if s.id.startswith("human_"))


def test_next_ten_successful_runs_only_and_reopen_is_idempotent(tmp_path):
    run(tmp_path, "old", frames=10)
    arm(tmp_path)
    assert register_run(tmp_path, "old")["runs"] == []
    for i in range(11):
        run(tmp_path, f"new{i}", frames=10)
        register_run(tmp_path, f"new{i}")
        register_run(tmp_path, f"new{i}")
    state = arm(tmp_path)
    assert len(state["runs"]) == 10
    assert state["runs"] == [f"new{i}" for i in range(10)]
    assert json.loads((tmp_path / "new0/meta.json").read_text())["training_run"]
    assert "training_run" not in json.loads((tmp_path / "new10/meta.json").read_text())


def test_middle_crop_is_five_seconds_and_skeleton_only(tmp_path):
    folder = run(tmp_path, "clip", frames=240)
    tracks, crop = middle_sequences(folder)
    assert crop["start_s"] == pytest.approx(1.5)
    assert crop["end_s"] == pytest.approx(6.5)
    assert len(tracks) > 0
    for track in tracks.values():
        assert track["sequence"].shape == (100, 51)
        assert np.isfinite(track["sequence"]).all()
        assert set(np.unique(track["sequence"][:, 34:])) <= {0, 1}


def test_short_clip_is_flagged_not_padded(tmp_path):
    folder = run(tmp_path, "short", frames=60)
    with pytest.raises(ValueError, match="at least 5s"):
        middle_sequences(folder)
    arm(tmp_path)  # Existing clip cannot accidentally enter the future cohort.
    folder = run(tmp_path, "newshort", frames=60)
    register_run(tmp_path, folder.name)
    state = refresh(tmp_path, train=False, export=False)
    assert not state["lstm_current"]
    assert any("at least 5s" in p for p in state["problems"])
    state = replace_run(tmp_path, folder.name)
    assert state["runs"] == []
    assert (folder / "tracking.jsonl").exists()
    assert register_run(tmp_path, folder.name)["runs"] == []


def records():
    rng = np.random.default_rng(123)
    x = rng.normal(size=(100, 35))
    y = (x[:, :8].sum(1) > 0).astype(int)
    return [{"run": f"r{i//10}", "label": int(y[i]), "values": dict(zip(PREDICTOR_IDS, x[i]))}
            for i in range(len(x))]


def test_lasso_returns_five_nonzero_ranked_weights_and_finite_scores():
    rows = records()
    model = fit_lasso(rows)
    assert len(model["feature_ids"]) == 5
    assert np.all(np.abs(model["ranking_coefficients"]) > 0)
    assert sorted(np.abs(model["ranking_coefficients"]), reverse=True) == list(np.abs(model["ranking_coefficients"]))
    assert set(model["feature_ids"]) <= set(PREDICTOR_IDS)
    result = score({"values": {}}, model)
    assert np.isfinite(result["score"])
    assert len(result["contributions"]) == 5


def test_lasso_refuses_to_invent_five_features():
    rows = records()
    for row in rows:
        row["values"] = {fid: 1.0 for fid in PREDICTOR_IDS}
    with pytest.raises(ValueError, match="fewer than five nonzero"):
        fit_lasso(rows)


def test_torch_trains_and_checkpoint_reloads(tmp_path):
    import torch
    rng = np.random.default_rng(2)
    samples = [{"sequence": rng.normal(size=(100,51)).astype("float32"), "label": i%2,
                "track_id": i, "run": f"r{i}"} for i in range(10)]
    report = train_lstm(samples, tmp_path, ["example"]*5, "test", epochs=2)
    assert len(report["runs"]) == 10
    assert np.isfinite(report["final_training_loss"])
    checkpoint = torch.load(tmp_path / "skeleton_lstm.pt", weights_only=True)
    model = make_model()
    model.load_state_dict(checkpoint["state_dict"])
    assert model(torch.zeros((2,100,51))).shape == (2,)
    data = np.load(tmp_path / "supervised_skeleton_sequences.npz")
    assert data["x"].shape == (10,100,51)


def test_training_starts_rolling_before_ten_and_discovers_added_run_folders(tmp_path, monkeypatch):
    import ddc.training as training
    arm(tmp_path)
    for i in range(2):
        run(tmp_path, f"clip{i}")
    def selected(rows):
        return {"feature_ids": PREDICTOR_IDS[:5], "ranking_coefficients": [.2]*5,
                "coefficient": [.2]*5, "intercept": .5, "medians": [0.]*5,
                "mean": [0.]*5, "std": [1.]*5, "runs": sorted({r["run"] for r in rows})}
    monkeypatch.setattr(training, "fit_lasso", selected)
    state = refresh(tmp_path, epochs=1, export=False)
    assert state["status"] == "rolling_trained"
    assert state["runs"] == ["clip0", "clip1"]
    assert state["lstm_runs"] == ["clip0", "clip1"]
    assert (tmp_path / "training/skeleton_lstm.pt").exists()
    # An unreviewed third run updates statistics without blocking or repeating training.
    third = run(tmp_path, "clip2")
    (third / "labels.csv").unlink()
    def repeat(*args, **kwargs):
        pytest.fail("an unreviewed new run must not refit unchanged supervised inputs")
    monkeypatch.setattr(training, "train_lstm", repeat)
    state = refresh(tmp_path, epochs=1, export=False)
    assert state["runs"] == ["clip0", "clip1", "clip2"]
    assert state["lstm_current"]
    payload = json.loads((tmp_path / "training_status.json").read_text())
    assert payload["lstm"]["samples"] == 16


def test_excel_label_import_is_atomic_and_never_uses_predictions(tmp_path, monkeypatch):
    # Excel is read with openpyxl; workbook authoring remains in the artifact builder.
    arm(tmp_path)
    folder = run(tmp_path, "clip")
    register_run(tmp_path, folder.name)
    (folder / "labels.csv").unlink()
    class Sheet:
        def __getitem__(self, key):
            return type("Cell", (), {"value": {"B2": "clip", "B3": signature(folder)}[key]})()
        def cell(self, row, col):
            return type("Cell", (), {"value": ["Track id", "Frames present", "Predicted label", "LASSO score", "Manual label", "Review status"][col-1]})()
        def iter_rows(self, **kwargs):
            return iter([(0,180,"dancer",1.0,"dancer","Reviewed"),
                         *[(i,180,"dancer",1.0,None,"Unreviewed") for i in range(1,8)]])
    class Book:
        sheetnames = ["Run 01"]
        def __getitem__(self, key): return Sheet()
        def close(self): pass
    import openpyxl
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: Book())
    workbook = tmp_path / "review.xlsx"
    workbook.touch()
    assert import_reviews(tmp_path, workbook) == 1
    from ddc.dancer_features import load_labels
    assert load_labels(folder / "labels.csv") == {0: 1}


def test_complete_cohort_trains_on_every_video_and_does_not_repeat_fit(tmp_path, monkeypatch):
    import ddc.training as training
    arm(tmp_path)
    for i in range(10):
        folder = run(tmp_path, f"clip{i}", frames=160)
        register_run(tmp_path, folder.name)
    # Selection numerics are tested above; this test isolates the cohort/training gate.
    def selected(rows):
        return {"feature_ids": PREDICTOR_IDS[:5], "ranking_coefficients": [.2]*5,
                "coefficient": [.2]*5, "intercept": .5, "medians": [0.]*5,
                "mean": [0.]*5, "std": [1.]*5, "runs": sorted({r["run"] for r in rows})}
    monkeypatch.setattr(training, "fit_lasso", selected)
    state = refresh(tmp_path, epochs=1, export=False)
    assert state["status"] == "trained"
    report = json.loads((tmp_path / "training/lstm_report.json").read_text())
    assert len(report["runs"]) == 10
    assert report["samples"] == 80
    def repeat(*args, **kwargs):
        pytest.fail("unchanged training data must not retrain")
    monkeypatch.setattr(training, "train_lstm", repeat)
    assert refresh(tmp_path, epochs=1, export=False)["status"] == "trained"
    # A failed new fit must not leave the old checkpoint marked current.
    label_path = tmp_path / "clip0/labels.csv"
    label_path.write_text(label_path.read_text().replace("0,dancer", "0,non_dancer"))
    def missing(*args, **kwargs):
        raise ImportError("test missing torch")
    monkeypatch.setattr(training, "train_lstm", missing)
    state = refresh(tmp_path, epochs=1, export=False)
    assert state["status"] == "training_failed"
    assert not state["lstm_current"]


def test_saved_review_updates_workbook_validation_labels_and_keeps_original_predictions(tmp_path, monkeypatch):
    import ddc.training as training
    from ddc.feature_registry import validate_checklist
    arm(tmp_path)
    folder = run(tmp_path, "clip")
    register_run(tmp_path, folder.name)
    predicted = [{"track_id": i, "predicted_label": "dancer" if i < 4 else "non_dancer",
                  "dancer_probability": .8 if i < 4 else .2} for i in range(8)]
    training.record_screening(tmp_path, folder.name, predicted, [])
    exports = []
    monkeypatch.setattr(training, "export_workbook", lambda root, payload: exports.append(payload))
    outcome = validate_checklist(tmp_path, "clip", [1,2,3,4], predicted)
    assert outcome["validation"]["corrected_track_ids"] == [0,4]
    from openpyxl import load_workbook
    legacy = load_workbook(tmp_path / "dancer_feature_registry.xlsx", read_only=True)
    assert legacy["Labels-clip"]["C5"].value == "non_dancer"
    assert not any(name.startswith("D1-") for name in legacy.sheetnames)
    legacy.close()
    rows = exports[-1]["runs"][0]["tracks"]
    assert rows[0]["manual_label"] == "non_dancer"
    assert rows[0]["prediction"]["predicted_label"] == "dancer"
    assert rows[0]["validation_status"] == "Corrected"
    assert rows[4]["manual_label"] == "dancer"
    assert rows[4]["review_action"] == "saved_review"
    # Refits cannot change the predictions against which those labels were reviewed.
    state = refresh(tmp_path, train=False, import_excel=False, export=False)
    assert state["reviewed_tracks"] == 8


def test_removed_run_does_not_block_training_registry_refresh(tmp_path):
    arm(tmp_path)
    folder = run(tmp_path, "removed", frames=30)
    register_run(tmp_path, folder.name)
    # Simulate tracking removed from its run without deleting any test files.
    (folder / "tracking.jsonl").rename(tmp_path / "retained_tracking.jsonl")
    state = refresh(tmp_path, train=False, import_excel=False, export=False)
    assert state["runs"] == []
    assert state["missing_runs"] == ["removed"]


def test_dancer_export_failure_does_not_stop_training_label_export(tmp_path, monkeypatch):
    import ddc.training as training
    import ddc.registry_workbook as registry_workbook
    from ddc.feature_registry import validate_checklist
    arm(tmp_path)
    folder = run(tmp_path, "clip", frames=160)
    register_run(tmp_path, folder.name)
    exports = []
    monkeypatch.setattr(training, "export_workbook", lambda root, payload: exports.append(payload))
    def unavailable(*args, **kwargs):
        raise PermissionError("dancer workbook is open")
    monkeypatch.setattr(registry_workbook, "write_registry_workbook", unavailable)
    outcome = validate_checklist(tmp_path, "clip", [0,1,2,3], [])
    assert "dancer workbook is open" in outcome["registry_note"]
    assert exports[-1]["runs"][0]["tracks"][0]["manual_label"] == "dancer"
    assert exports[-1]["runs"][0]["tracks"][4]["manual_label"] == "non_dancer"


def test_screening_does_not_persist_individual_predictor_statistics(tmp_path):
    from ddc.training import record_screening
    folder = run(tmp_path, "clip", frames=30)
    record_screening(tmp_path, "clip", [{"track_id": 0, "predicted_label": "dancer",
                      "dancer_probability": .8, "top_predictors": [{"raw_value": 12}]}], [])
    saved = json.loads((folder / "dancer_screening.json").read_text())
    assert "top_predictors" not in saved["predictions"][0]
