import json

from ddc.dancer_features import write_synthetic_demo
from ddc.feature_registry import build_registry, predict_run, validate_checklist


def _make_run(root, name, *, labels):
    run = root / name
    run.mkdir()
    label_path = run / "labels.csv" if labels else root / f"{name}-source-labels.csv"
    write_synthetic_demo(run / "tracking.jsonl", label_path, frames=40)
    return run


def test_registry_scores_with_only_the_workbook_top_five_features(tmp_path):
    _make_run(tmp_path, "training", labels=True)
    _make_run(tmp_path, "candidate", labels=False)

    model, predictions = predict_run(tmp_path, "candidate")

    assert len(model["feature_ids"]) == 5
    assert model["feature_ids"] == [item["feature_id"] for item in model["selected_features"]]
    assert all(len(item["top_predictors"]) == 5 for item in predictions)
    assert all(
        [feature["feature_id"] for feature in item["top_predictors"]] == model["feature_ids"]
        for item in predictions
    )


def test_review_validation_records_corrections_and_preserves_original_prediction(tmp_path):
    _make_run(tmp_path, "training", labels=True)
    candidate = _make_run(tmp_path, "candidate", labels=False)
    _, predictions = predict_run(tmp_path, "candidate")
    predicted_dancers = [item["track_id"] for item in predictions
                         if item["predicted_label"] == "dancer"]

    unchanged = validate_checklist(tmp_path, "candidate", predicted_dancers, predictions)
    assert unchanged["validation"]["status"] == "validated_correct"
    assert unchanged["validation"]["corrected_track_ids"] == []

    changed_selection = [track for track in predicted_dancers if track != 0]
    if 0 not in predicted_dancers:
        changed_selection.append(0)
    corrected = validate_checklist(tmp_path, "candidate", changed_selection, predictions)
    assert corrected["validation"]["status"] == "corrected"
    assert 0 in corrected["validation"]["corrected_track_ids"]

    validation = json.loads((candidate / "prediction_validation.json").read_text())
    assert validation["predictor_labels"] == {
        str(item["track_id"]): item["predicted_label"] for item in predictions
    }
    rebuilt = build_registry(tmp_path)
    candidate_detail = next(run for run in rebuilt["runs_detail"] if run["run"] == "candidate")
    assert {str(track["track_id"]): track["predicted_label"] for track in candidate_detail["tracks"]} == validation["predictor_labels"]


def test_scored_run_is_held_out_of_training(tmp_path):
    _make_run(tmp_path, "training", labels=True)
    _make_run(tmp_path, "candidate", labels=True)  # already reviewed

    model, _ = predict_run(tmp_path, "candidate")

    assert model["held_out_run"] == "candidate"
    assert model["runs"] == ["training"]


def test_scoring_needs_labels_from_another_run(tmp_path):
    import pytest
    _make_run(tmp_path, "only", labels=True)
    with pytest.raises(ValueError):
        predict_run(tmp_path, "only")


def test_camera_context_features_are_never_predictors(tmp_path):
    from ddc.feature_inventory import ELIGIBLE_FEATURE_IDS
    assert not [f for f in ELIGIBLE_FEATURE_IDS if f.startswith("context_")]
    _make_run(tmp_path, "training", labels=True)
    _make_run(tmp_path, "candidate", labels=False)
    model, _ = predict_run(tmp_path, "candidate")
    assert not [f for f in model["feature_ids"] if f.startswith("context_")]


def test_workbook_is_written_with_summary_labels_and_dancer_sheets(tmp_path):
    from openpyxl import load_workbook
    from ddc.feature_registry import rebuild_registry_workbook
    _make_run(tmp_path, "training", labels=True)
    _make_run(tmp_path, "candidate", labels=False)

    path = rebuild_registry_workbook(tmp_path)

    wb = load_workbook(path)
    assert wb.sheetnames[:3] == ["Registry Summary", "Feature Inventory Summary", "Feature Catalog"]
    assert "Labels-training" in wb.sheetnames and "Labels-candidate" in wb.sheetnames
    assert any(name.startswith("D1-training") for name in wb.sheetnames)
    summary = wb["Registry Summary"]
    assert [summary.cell(row=r, column=2).value for r in range(11, 16)]  # five predictors
    assert wb["Labels-training"]["A5"].value == 0


def test_workbook_failure_is_a_runtime_error(tmp_path):
    import pytest
    from ddc.feature_registry import rebuild_registry_workbook
    _make_run(tmp_path, "only", labels=False)  # no labels anywhere -> cannot fit
    with pytest.raises(RuntimeError, match="workbook builder failed"):
        rebuild_registry_workbook(tmp_path)


def test_missing_openpyxl_is_a_runtime_error(tmp_path, monkeypatch):
    import builtins
    import pytest
    from ddc.feature_registry import rebuild_registry_workbook
    _make_run(tmp_path, "training", labels=True)
    _make_run(tmp_path, "candidate", labels=False)
    real_import = builtins.__import__

    def fake(name, *a, **k):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("No module named 'openpyxl'")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake)
    monkeypatch.delitem(__import__("sys").modules, "ddc.registry_workbook", raising=False)
    with pytest.raises(RuntimeError, match="pip install"):
        rebuild_registry_workbook(tmp_path)
